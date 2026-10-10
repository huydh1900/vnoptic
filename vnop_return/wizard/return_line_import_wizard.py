# -*- coding: utf-8 -*-
import base64
import io
import logging

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from odoo import fields, models, _
from odoo.exceptions import UserError
from odoo.tools import float_compare

_logger = logging.getLogger(__name__)

TEMPLATE_COLUMNS = [
    ("default_code", "Mã sản phẩm", "Mã nội bộ (default_code), required"),
    ("quantity", "Số lượng", "Số, required"),
]

BASE_FONT_NAME = "Times New Roman"
BASE_FONT_SIZE = 14
HEADER_FONT = Font(name=BASE_FONT_NAME, bold=True, color="FFFFFF", size=BASE_FONT_SIZE)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HINT_FONT = Font(name=BASE_FONT_NAME, italic=True, color="808080", size=BASE_FONT_SIZE)
REQUIRED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
LABEL_FONT = Font(name=BASE_FONT_NAME, bold=True, size=BASE_FONT_SIZE)
DATA_ALIGNMENT = Alignment(horizontal="left")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


class ReturnLineImportWizard(models.TransientModel):
    _name = "vnop.return.line.import.wizard"
    _description = "Import dòng trả hàng từ Excel"

    request_id = fields.Many2one("vnop.return.request", required=True, readonly=True)
    return_type = fields.Selection(related="request_id.return_type")
    defect_group = fields.Many2one(
        "vnop.return.defect.group", string="Nhóm lỗi", required=True,
        domain="[('return_type', '=', return_type)]")
    reason_id = fields.Many2one(
        "vnop.return.reason", string="Lý do", required=True,
        domain="[('defect_group', '=', defect_group)]")
    file_data = fields.Binary(string="File Excel")
    file_name = fields.Char()

    state = fields.Selection([
        ("upload", "Upload"),
        ("validated", "Kiểm thử"),
        ("done", "Hoàn tất"),
    ], default="upload")
    validated_count = fields.Integer(string="Dòng hợp lệ", readonly=True)
    imported_count = fields.Integer(readonly=True)
    error_text = fields.Text(readonly=True)

    def action_validate(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Vui lòng upload file Excel."))
        raw = base64.b64decode(self.file_data)
        rows, errors = self._parse_excel(raw)
        if not errors:
            _vals, errors = self._prepare_line_vals(rows)
        self.write({
            "state": "validated",
            "validated_count": len(rows) - len(errors) if not errors else 0,
            "error_text": "\n".join(errors) if errors else False,
        })
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Vui lòng upload file Excel."))
        raw = base64.b64decode(self.file_data)
        rows, errors = self._parse_excel(raw)
        if errors:
            self.write({
                "state": "done",
                "imported_count": 0,
                "error_text": "\n".join(errors),
            })
            return self._reopen()

        line_vals_list, import_errors = self._prepare_line_vals(rows)
        if import_errors:
            self.write({
                "state": "done",
                "imported_count": 0,
                "error_text": "\n".join(import_errors),
            })
            return self._reopen()

        if not line_vals_list:
            raise UserError(_("File không có dữ liệu hợp lệ."))

        self.env["vnop.return.request.line"].create(line_vals_list)
        self.write({
            "state": "done",
            "imported_count": len(line_vals_list),
        })
        return self._reopen()

    def action_back_upload(self):
        self.ensure_one()
        self.write({
            "state": "upload",
            "validated_count": 0,
            "imported_count": 0,
            "error_text": False,
        })
        return self._reopen()

    def _parse_excel(self, file_bytes):
        errors = []
        rows = []
        try:
            wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception as e:
            return [], [_("Không đọc được file Excel: %s") % str(e)]

        ws = wb.active
        if ws is None:
            wb.close()
            return [], [_("File Excel không có sheet nào.")]

        header_row = []
        for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=False), []):
            header_row.append(str(cell.value).strip() if cell.value else "")

        col_keys = [c[0] for c in TEMPLATE_COLUMNS]
        col_mapping = {}
        for col_idx, val in enumerate(header_row):
            if val in col_keys:
                col_mapping[val] = col_idx

        missing = [k for k in col_keys if k not in col_mapping]
        if missing:
            wb.close()
            return [], [_("Thiếu cột bắt buộc: %s") % ", ".join(missing)]

        for row_idx, row_cells in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            if row_cells is None:
                continue
            if all(c is None or str(c).strip() == "" for c in row_cells):
                continue
            row_data = {"_row": row_idx}
            for key, ci in col_mapping.items():
                row_data[key] = row_cells[ci] if ci < len(row_cells) else None
            rows.append(row_data)

        wb.close()
        if not rows:
            errors.append(_("File không có dữ liệu (data bắt đầu từ row 3)."))
        return rows, errors

    def _prepare_line_vals(self, rows):
        errors = []
        request = self.request_id
        so = request.sale_order_id

        # Gộp các dòng cùng default_code
        merged = {}
        for row in rows:
            code = self._normalize_code(row.get("default_code"))
            if not code:
                errors.append(_("Dòng %d: Thiếu mã sản phẩm.") % row["_row"])
                continue
            try:
                qty = float(row.get("quantity") or 0)
                if qty <= 0:
                    raise ValueError
            except (ValueError, TypeError):
                errors.append(_("Dòng %d: Số lượng không hợp lệ.") % row["_row"])
                continue
            if code in merged:
                merged[code]["quantity"] += qty
                merged[code]["_rows"].append(row["_row"])
            else:
                merged[code] = {"quantity": qty, "_rows": [row["_row"]]}

        if errors:
            return [], errors

        # Tìm sản phẩm
        codes = list(merged.keys())
        products = self.env["product.product"].search([("default_code", "in", codes)])
        product_by_code = {p.default_code: p for p in products if p.default_code}

        # Allowed products (nếu có đơn gốc)
        allowed = request.allowed_product_ids if so else False

        # Tính SL đã trả từ các phiếu active khác
        already_returned = {}
        if so:
            existing_lines = self.env["vnop.return.request.line"].search([
                ("request_id.sale_order_id", "=", so.id),
                ("request_id.state", "not in", ("cancel", "refused")),
                ("request_id", "!=", request.id),
            ])
            for el in existing_lines:
                pid = el.product_id.id
                already_returned[pid] = already_returned.get(pid, 0.0) + el.quantity
            # Cộng thêm SL đã có trên phiếu hiện tại
            for el in request.line_ids:
                pid = el.product_id.id
                already_returned[pid] = already_returned.get(pid, 0.0) + el.quantity

        # qty_delivered theo product trên đơn gốc (cộng tất cả dòng cùng SP)
        delivered_by_product = {}
        if so:
            for sol in so.order_line:
                pid = sol.product_id.id
                delivered_by_product[pid] = delivered_by_product.get(pid, 0.0) + sol.qty_delivered

        uom_precision = self.env["decimal.precision"].precision_get("Product Unit of Measure")
        vals_list = []

        for code, data in merged.items():
            row_label = ", ".join(str(r) for r in data["_rows"])
            product = product_by_code.get(code)
            if not product:
                errors.append(_("Dòng %s: Không tìm thấy SP với mã '%s'.") % (row_label, code))
                continue

            if allowed and product not in allowed:
                errors.append(_("Dòng %s: SP '%s' không thuộc đơn hàng gốc.") % (row_label, code))
                continue

            qty = data["quantity"]

            # Check SL trả vượt qty_delivered
            if so:
                total_after = already_returned.get(product.id, 0.0) + qty
                max_qty = delivered_by_product.get(product.id, 0.0)
                if float_compare(total_after, max_qty, precision_digits=uom_precision) > 0:
                    errors.append(_(
                        "Dòng %s: SP '%s' — tổng SL trả (%.2f) vượt SL đã giao (%.2f)."
                    ) % (row_label, code, total_after, max_qty))
                    continue

            # Lấy UoM từ dòng đơn gốc hoặc SP
            sol = so.order_line.filtered(
                lambda ol, p=product: ol.product_id == p)[:1] if so else False
            uom = sol.product_uom if sol else product.uom_id

            vals_list.append({
                "request_id": request.id,
                "product_id": product.id,
                "product_uom_id": uom.id,
                "quantity": qty,
                "defect_group": self.defect_group.id,
                "reason_id": self.reason_id.id,
            })

        return vals_list, errors

    @staticmethod
    def _normalize_code(raw):
        if raw is None:
            return ""
        if isinstance(raw, float):
            if raw == int(raw):
                return str(int(raw)).strip()
            return str(raw).strip()
        return str(raw).strip()

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    @staticmethod
    def generate_template():
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"

        for col_idx, (key, label, hint) in enumerate(TEMPLATE_COLUMNS, 1):
            col_letter = get_column_letter(col_idx)

            cell = ws.cell(row=1, column=col_idx, value=key)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = DATA_ALIGNMENT
            cell.border = THIN_BORDER

            cell_label = ws.cell(row=2, column=col_idx, value=label)
            cell_label.font = LABEL_FONT
            cell_label.alignment = DATA_ALIGNMENT
            cell_label.border = THIN_BORDER
            if "required" in hint.lower():
                cell_label.fill = REQUIRED_FILL

            max_len = max(len(key), len(label))
            ws.column_dimensions[col_letter].width = max(max_len * 1.4 + 4, 18)

        text_col_indices = [
            i for i, (k, _l, _h) in enumerate(TEMPLATE_COLUMNS, 1) if k == "default_code"
        ]
        for row in range(3, 1003):
            for col_idx in range(1, len(TEMPLATE_COLUMNS) + 1):
                c = ws.cell(row=row, column=col_idx)
                c.font = Font(name=BASE_FONT_NAME, size=BASE_FONT_SIZE)
                c.alignment = DATA_ALIGNMENT
                if col_idx in text_col_indices:
                    c.number_format = numbers.FORMAT_TEXT

        ws.freeze_panes = "A3"

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
